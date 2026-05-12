"""文档解析模块：支持 PDF、DOCX、Markdown、HTML、CSV、TXT 等多种格式。"""
from __future__ import annotations

import csv as csv_module
import re
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path


# ============ 文本提取 ============

class _HTMLTextExtractor(HTMLParser):
    """从 HTML 中提取纯文本。"""

    def __init__(self):
        super().__init__()
        self._parts: list[str] = []
        self._skip_tags = {"script", "style", "noscript", "code", "pre"}

    def handle_data(self, data: str):
        if self.get_starttag_text() and self._in_skip_tag():
            return
        text = data.strip()
        if text:
            self._parts.append(text)

    def _in_skip_tag(self) -> bool:
        for tag in self._skip_tags:
            if tag in (self.get_starttag_text() or "").lower():
                return True
        return False

    def get_text(self) -> str:
        return "\n".join(self._parts)


def parse_pdf(file_bytes: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    texts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n".join(texts)


def parse_docx(file_bytes: bytes) -> str:
    from io import BytesIO as _BytesIO
    from docx import Document

    doc = Document(_BytesIO(file_bytes))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    # 也提取表格中的文本
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_html(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    extractor = _HTMLTextExtractor()
    extractor.feed(text)
    return extractor.get_text()


def parse_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    reader = csv_module.reader(StringIO(text))
    rows: list[str] = []
    headers = next(reader, None)
    if headers:
        rows.append(" | ".join(h.strip() for h in headers))
    for row in reader:
        if row:
            rows.append(" | ".join(c.strip() for c in row))
    return "\n".join(rows)


def parse_markdown(text: str) -> str:
    import re as _re

    text = _re.sub(r"```[\s\S]*?```", "", text)
    text = _re.sub(r"`[^`]+`", "", text)
    text = _re.sub(r"^#{1,6}\s+", "", text, flags=_re.MULTILINE)
    text = _re.sub(r"!\[.*?\]\(.*?\)", "", text, flags=_re.DOTALL)
    text = _re.sub(r"\[([^\]]*)\]\(.*?\)", r"\1", text, flags=_re.DOTALL)
    text = _re.sub(r"\*{1,3}([^*]+)\*{1,3}", r"\1", text)
    text = _re.sub(r"^[\s]*[-*+]\s+", "", text, flags=_re.MULTILINE)
    text = _re.sub(r"^[\s]*\d+\.\s+", "", text, flags=_re.MULTILINE)
    text = _re.sub(r"^[-*_]{3,}\s*$", "", text, flags=_re.MULTILINE)
    text = _re.sub(r"^>\s+", "", text, flags=_re.MULTILINE)
    text = _re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ============ 文本分块 ============

def chunk_text(
    text: str,
    chunk_size: int = 500,
    chunk_overlap: int = 100,
) -> list[str]:
    """按段落和字符数分块，保持段落完整性。"""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for para in paragraphs:
        if current_len + len(para) <= chunk_size:
            current.append(para)
            current_len += len(para)
        else:
            if current:
                chunks.append("\n\n".join(current))
            if len(para) > chunk_size:
                step = max(50, chunk_size - chunk_overlap)
                for i in range(0, len(para), step):
                    chunks.append(para[i:i + chunk_size])
                current = []
                current_len = 0
            else:
                current = [para]
                current_len = len(para)

    if current:
        chunks.append("\n\n".join(current))

    return chunks


# ============ 主入口 ============

def extract_text(filename: str, file_bytes: bytes) -> str:
    """提取文件中的纯文本（不切块），用于预览。"""
    ext = Path(filename).suffix.lower()
    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".html": parse_html,
        ".htm": parse_html,
        ".csv": parse_csv,
    }

    if ext in parsers:
        return parsers[ext](file_bytes)

    text = file_bytes.decode("utf-8", errors="replace")
    if ext in (".md", ".markdown"):
        return parse_markdown(text)
    return text.strip()


def parse_file(
    filename: str,
    file_bytes: bytes,
    chunk_size: int = 500,
    chunk_overlap: int = 100,
) -> list[str]:
    """解析文件并返回 chunk 列表。"""
    text = extract_text(filename, file_bytes)
    if not text:
        return []
    return chunk_text(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
